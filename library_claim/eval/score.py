"""Ground-truth template and scorer for the pass bars.

    python -m library_claim.cli gt-template ground_truth.xlsx
    (fill it in by hand: count, read titles, tape-measure, check prices)
    python -m library_claim.cli score sweeps/<id> ground_truth.xlsx

Matching rules (so the score cannot be gamed by the matcher):
* A ground-truth book matches a packet book through the ``packet_id`` column
  if you filled it in, otherwise by title (fuzzy >= 0.85 on normalised title)
  among *identified* packet books, one-to-one.
* "Confidently wrong" = packet lines with status ``identified`` whose title
  matches no ground-truth title at all.
* Dimension and price bars are scored as mean absolute percentage error over
  the sample (assumption; the brief does not say mean or per-item), and the
  share of the sample inside the bar is reported alongside.
"""

from __future__ import annotations

import json
from difflib import SequenceMatcher
from pathlib import Path
from statistics import mean, median

from ..vision.tracker import norm_text

SHEETS = {
    "summary": ["field", "value", "notes"],
    "books": ["gt_id", "shelf", "position", "legible_by_human(Y/N)", "title", "author", "packet_id(optional)"],
    "spine_sample": ["gt_id", "title_or_description", "height_cm", "thickness_cm", "packet_id(optional)"],
    "price_sample": ["title", "type(new/used)", "retailer", "url", "checked_on", "price", "currency", "packet_id(optional)"],
    "items": ["gt_item", "category", "description", "packet_id(optional)"],
    "room": ["field", "value_m", "notes"],
}
SUMMARY_ROWS = [["true_book_count", "", "every book you count by hand, including unreadable and stacked ones"],
                ["country", "US", ""], ["currency", "USD", ""]]
ROOM_ROWS = [["length_m", "", "tape"], ["width_m", "", "tape"], ["height_m", "", "tape (floor to ceiling)"],
             ["floor_area_m2_override", "", "only if the room is not a rectangle"],
             ["wall_area_m2_override", "", "only if not 2*(L+W)*H (gross, no door/window deduction)"]]
BARS = {"book_count": 0.05, "title_correct": 0.70, "title_wrong": 0.03, "spine": 0.15, "price": 0.25, "items": 0.80,
        "floor": 0.10, "wall": 0.15, "time_s": 300}


def write_template(path: Path) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for r in {"summary": SUMMARY_ROWS, "room": ROOM_ROWS}.get(name, []):
            ws.append(r)
    wb.save(path)


def _rows(wb, name):
    ws = wb[name]
    hdr = [c.value for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in r):
            out.append(dict(zip(hdr, r)))
    return out


def _sim(a, b) -> float:
    a, b = norm_text(str(a or "")), norm_text(str(b or ""))
    return SequenceMatcher(None, a, b).ratio() if a and b else 0.0


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def match_titles(gt_books, packet_books, threshold=0.85):
    """One-to-one title matching. Returns {gt_id: packet_id}."""
    by_id = {b["id"]: b for b in packet_books}
    used, out = set(), {}
    for g in gt_books:
        pid = g.get("packet_id(optional)")
        if pid and pid in by_id:
            out[g["gt_id"]] = pid
            used.add(pid)
    cands = []
    for g in gt_books:
        if g["gt_id"] in out or not g.get("title"):
            continue
        for b in packet_books:
            if b["status"] == "identified" and b["id"] not in used:
                s = _sim(g["title"], b["title"].split(":")[0])
                if s >= threshold:
                    cands.append((s, g["gt_id"], b["id"]))
    for s, gid, pid in sorted(cands, reverse=True):
        if gid not in out and pid not in used:
            out[gid] = pid
            used.add(pid)
    return out


def score(sweep_dir: Path, gt_path: Path) -> dict:
    from openpyxl import load_workbook
    p = json.loads((sweep_dir / "claim_packet.json").read_text(encoding="utf-8"))
    wb = load_workbook(gt_path, data_only=True)
    summ = {r["field"]: r["value"] for r in _rows(wb, "summary")}
    gt_books = _rows(wb, "books")
    res: dict = {"sweep": p["sweep"]["id"], "bars": {}}

    true_n = _num(summ.get("true_book_count")) or len(gt_books)
    got_n = p["totals"]["book_count"]
    err = abs(got_n - true_n) / true_n if true_n else None
    res["bars"]["book_count"] = {"true": true_n, "system": got_n, "error": err, "bar": "<= 5%",
                                 "pass": err is not None and err <= BARS["book_count"]}

    legible = [g for g in gt_books if str(g.get("legible_by_human(Y/N)", "")).strip().upper().startswith("Y")]
    m = match_titles(legible, p["books"])
    by_id = {b["id"]: b for b in p["books"]}
    correct = sum(1 for gid, pid in m.items() if by_id[pid]["status"] == "identified"
                  and _sim(next(g["title"] for g in legible if g["gt_id"] == gid), by_id[pid]["title"].split(":")[0]) >= 0.85)
    all_titles = [g.get("title") for g in gt_books if g.get("title")]
    wrong = [b for b in p["books"] if b["status"] == "identified"
             and not any(_sim(t, b["title"].split(":")[0]) >= 0.85 for t in all_titles)]
    n_leg = len(legible)
    res["bars"]["title_identification"] = {
        "human_legible": n_leg, "correct": correct, "correct_rate": correct / n_leg if n_leg else None,
        "confidently_wrong": len(wrong), "wrong_rate": len(wrong) / n_leg if n_leg else None,
        "wrong_lines": [{"id": b["id"], "title": b["title"], "spine_text": b["spine_text"], "frame_ref": b["frame_ref"]} for b in wrong],
        "bar": ">= 70% correct and <= 3% confidently wrong",
        "pass": bool(n_leg) and correct / n_leg >= BARS["title_correct"] and len(wrong) / n_leg <= BARS["title_wrong"]}

    sample = _rows(wb, "spine_sample")
    sm = match_titles([{**r, "title": r.get("title_or_description")} for r in sample], p["books"], 0.85)
    errs_h, errs_t, lines = [], [], []
    for r in sample:
        pid = r.get("packet_id(optional)") or sm.get(r["gt_id"])
        b = by_id.get(pid)
        hh, ht = _num(r.get("height_cm")), _num(r.get("thickness_cm"))
        line = {"gt_id": r["gt_id"], "packet_id": pid, "true_h": hh, "true_t": ht,
                "sys_h": b and b["spine_height_cm"], "sys_t": b and b["spine_thickness_cm"]}
        if b and hh and b["spine_height_cm"]:
            errs_h.append(abs(b["spine_height_cm"] - hh) / hh)
        if b and ht and b["spine_thickness_cm"]:
            errs_t.append(abs(b["spine_thickness_cm"] - ht) / ht)
        lines.append(line)
    unmatched = sum(1 for l in lines if l["sys_h"] is None)
    res["bars"]["spine_dimensions"] = {
        "sample": len(sample), "matched_with_value": len(errs_h), "missing_or_unmatched": unmatched,
        "height_mape": mean(errs_h) if errs_h else None, "thickness_mape": mean(errs_t) if errs_t else None,
        "height_within_15pct": sum(e <= .15 for e in errs_h) / len(sample) if sample else None,
        "thickness_within_15pct": sum(e <= .15 for e in errs_t) / len(sample) if sample else None,
        "lines": lines, "bar": "within 15% (MAPE over the hand-measured sample; missing values count as misses in the share)",
        "pass": bool(errs_h) and mean(errs_h) <= BARS["spine"] and bool(errs_t) and mean(errs_t) <= BARS["spine"]
                and unmatched == 0}

    ps = _rows(wb, "price_sample")
    perrs, plines = [], []
    pm = match_titles([{**r, "gt_id": i, "title": r.get("title")} for i, r in enumerate(ps)], p["books"], 0.85)
    for i, r in enumerate(ps):
        pid = r.get("packet_id(optional)") or pm.get(i)
        b = by_id.get(pid)
        kind = str(r.get("type(new/used)") or "new").lower()
        sysv = b and (b["used_value"]["amount"] if kind.startswith("u") else b["replacement_cost"]["amount"])
        hand = _num(r.get("price"))
        e = abs(sysv - hand) / hand if (sysv and hand) else None
        if e is not None:
            perrs.append(e)
        plines.append({"title": r.get("title"), "type": kind, "hand": hand, "system": sysv, "error": e,
                       "system_source": b and (b["used_value"] if kind.startswith("u") else b["replacement_cost"]).get("url")})
    res["bars"]["book_prices"] = {"sample": len(ps), "priced": len(perrs), "mape": mean(perrs) if perrs else None,
                                  "within_25pct": sum(e <= .25 for e in perrs) / len(ps) if ps else None, "lines": plines,
                                  "bar": "within 25% (MAPE), same source type",
                                  "pass": bool(perrs) and mean(perrs) <= BARS["price"] and len(perrs) == len(ps)}

    gi = _rows(wb, "items")
    found, ilines, used = 0, [], set()
    for g in gi:
        pid = g.get("packet_id(optional)")
        hit = next((i for i in p["items"] if i["id"] == pid), None) if pid else None
        if hit is None:
            cands = [i for i in p["items"] if i["id"] not in used and i["category"] == str(g.get("category", "")).lower()]
            cands.sort(key=lambda i: _sim(g.get("description"), i["description"]), reverse=True)
            hit = cands[0] if cands and _sim(g.get("description"), cands[0]["description"]) > 0.3 else None
        ok = bool(hit) and hit["category"] == str(g.get("category", "")).lower()
        if hit:
            used.add(hit["id"])
        found += ok
        ilines.append({"gt": g.get("description"), "category": g.get("category"), "packet_id": hit and hit["id"], "ok": ok})
    res["bars"]["non_book_items"] = {"true": len(gi), "found_correct_category": found, "rate": found / len(gi) if gi else None,
                                     "extra_system_items": len(p["items"]) - len(used), "lines": ilines, "bar": ">= 80%",
                                     "pass": bool(gi) and found / len(gi) >= BARS["items"]}

    room = {r["field"]: _num(r["value_m"]) for r in _rows(wb, "room")}
    L, W, H = room.get("length_m"), room.get("width_m"), room.get("height_m")
    floor = room.get("floor_area_m2_override") or (L * W if L and W else None)
    wall = room.get("wall_area_m2_override") or (2 * (L + W) * H if L and W and H else None)
    pr = p["room"]
    for key, true, sysv, bar in (("floor_area", floor, pr["floor_area_m2"], BARS["floor"]), ("wall_area", wall, pr["wall_area_m2"], BARS["wall"])):
        e = abs(sysv - true) / true if (true and sysv) else None
        res["bars"][key] = {"true_m2": true, "system_m2": sysv, "error": e, "bar": f"<= {bar:.0%}", "pass": e is not None and e <= bar}
    res["bars"]["room_dims"] = {"true": [L, W, H], "system": [pr["length_m"], pr["width_m"], pr["height_m"]]}
    tp = p["meta"].get("time_to_packet_s")
    res["bars"]["time_to_packet"] = {"seconds": tp, "bar": "< 300 s", "pass": tp is not None and tp < BARS["time_s"]}
    res["cost_usd"] = p["meta"].get("cost", {}).get("total_cost_usd")
    res["passed"] = sum(1 for v in res["bars"].values() if v.get("pass"))
    res["scored"] = sum(1 for v in res["bars"].values() if "pass" in v)
    return res


def write_results(res: dict, gt_path: Path, out_dir: Path) -> Path:
    from openpyxl import load_workbook
    wb = load_workbook(gt_path)
    if "results" in wb.sheetnames:
        del wb["results"]
    ws = wb.create_sheet("results")
    ws.append(["measure", "bar", "true", "system", "error / rate", "pass"])
    b = res["bars"]
    ws.append(["Book count", b["book_count"]["bar"], b["book_count"]["true"], b["book_count"]["system"], b["book_count"]["error"], b["book_count"]["pass"]])
    t = b["title_identification"]
    ws.append(["Title identification", t["bar"], t["human_legible"], f"{t['correct']} correct / {t['confidently_wrong']} wrong", f"{t['correct_rate']} / {t['wrong_rate']}", t["pass"]])
    s = b["spine_dimensions"]
    ws.append(["Spine dimensions", s["bar"], s["sample"], f"H MAPE {s['height_mape']}, T MAPE {s['thickness_mape']}", f"H in-bar {s['height_within_15pct']}, T in-bar {s['thickness_within_15pct']}", s["pass"]])
    pr = b["book_prices"]
    ws.append(["Book prices", pr["bar"], pr["sample"], pr["priced"], f"MAPE {pr['mape']}, in-bar {pr['within_25pct']}", pr["pass"]])
    it = b["non_book_items"]
    ws.append(["Non-book items", it["bar"], it["true"], it["found_correct_category"], it["rate"], it["pass"]])
    for k, label in (("floor_area", "Floor area"), ("wall_area", "Wall area")):
        ws.append([label, b[k]["bar"], b[k]["true_m2"], b[k]["system_m2"], b[k]["error"], b[k]["pass"]])
    ws.append(["Time to packet", b["time_to_packet"]["bar"], "", b["time_to_packet"]["seconds"], "", b["time_to_packet"]["pass"]])
    ws.append([])
    ws.append(["cost per sweep (USD)", res.get("cost_usd")])
    out = out_dir / "ground_truth_scored.xlsx"
    wb.save(out)
    (out_dir / "score.json").write_text(json.dumps(res, indent=2, default=str), encoding="utf-8")
    return out


__all__ = ["write_template", "score", "write_results", "median"]
