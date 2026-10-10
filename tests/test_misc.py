"""Totals, capture quality, director nudges, scorer, marker scale, live server smoke."""

import json

import cv2
import numpy as np

from library_claim.capture.quality import assess
from library_claim.live.agent import Director
from library_claim.measure.scale import detect_markers, plane_from_marker, quad_dims_cm, render_marker_png
from library_claim.packet.totals import compute_totals
from library_claim.schemas import Book, Item, ItemCost, Money, UsedMoney


def test_totals_from_lines_only():
    books = [
        Book(id="a", status="identified", replacement_cost=Money(amount=10.0), used_value=UsedMoney(amount=4.0)),
        Book(id="b", status="identified", replacement_cost=Money(amount=None), used_value=UsedMoney(amount=2.5)),
        Book(id="c", status="unidentified"),
        Book(id="d", status="needs_appraisal", replacement_cost=Money(amount=999.0)),   # never counted
        Book(id="e", status="identified"),
    ]
    items = [Item(id="i1", category="lamp", description="", status="range", replacement_cost=ItemCost(low=20, high=40)),
             Item(id="i2", category="art", description="", status="needs_appraisal")]
    t = compute_totals(books, items, "USD", 1.2)
    assert t.book_count == 5 and t.books_identified == 3 and t.books_unidentified == 1 and t.books_needs_appraisal == 1
    assert t.books_replacement_cost == 10.0 and t.books_used_value == 6.5
    assert t.items_replacement_cost_low == 20 and t.items_replacement_cost_high == 40
    assert t.excluded_from_totals == 4   # c, d, e, i2


def test_quality_flags():
    rng = np.random.default_rng(0)
    sharp = (rng.random((480, 640, 3)) * 255).astype(np.uint8)
    assert assess(sharp).issues == []
    blurred = cv2.GaussianBlur(sharp, (31, 31), 10)
    assert "blur" in assess(blurred).issues
    glare = sharp.copy()
    glare[:120] = 255
    assert "glare" in assess(glare).issues
    assert "dark" in assess((sharp * 0.05).astype(np.uint8)).issues
    assert "fast" in assess(sharp, motion_px=80).issues


def test_marker_scale_exact():
    mk = render_marker_png(0, 300)                  # 300 px marker + border
    canvas = np.full((900, 1200), 200, np.uint8)
    canvas[300:300 + mk.shape[0], 400:400 + mk.shape[1]] = mk
    hits = detect_markers(cv2.cvtColor(canvas, cv2.COLOR_GRAY2BGR), "f")
    assert hits
    P = plane_from_marker(hits[0], 15.0)            # 300 px == 15 cm -> 20 px / cm
    w, h = quad_dims_cm(P, np.float32([[0, 0], [200, 0], [200, 400], [0, 400]]))
    assert abs(w - 10.0) < 0.15 and abs(h - 20.0) < 0.3


def test_director_rate_limits_and_prioritises():
    d = Director()
    st = {"problems": ["blur"], "spines": 10, "legible": 4, "items": 1}
    n1 = d.next_notice(st, [], False, now=100.0)
    assert n1 and n1.startswith("NUDGE") and "steady" in n1
    assert d.next_notice(st, [], False, now=102.0) is None        # global gap
    ask = d.next_notice(st, [{"item_id": "I1", "description": "portrait"}], False, now=110.0)
    assert ask.startswith("ASK") and "I1" in ask


def test_scorer_on_demo_packet(tmp_path):
    from library_claim.eval.score import score, write_template
    from openpyxl import load_workbook
    packet = {"sweep": {"id": "s", "country": "US", "currency": "USD"},
              "books": [{"id": "B1", "status": "identified", "title": "The Hobbit", "spine_text": "THE HOBBIT", "frame_ref": "f",
                         "spine_height_cm": 20.0, "spine_thickness_cm": 2.0, "replacement_cost": {"amount": 15.0, "url": "u"},
                         "used_value": {"amount": 5.0, "url": "u"}},
                        {"id": "B2", "status": "identified", "title": "Wrong Book", "spine_text": "X", "frame_ref": "f",
                         "spine_height_cm": None, "spine_thickness_cm": None, "replacement_cost": {"amount": None}, "used_value": {"amount": None}},
                        {"id": "B3", "status": "unidentified", "title": "", "spine_text": "", "frame_ref": "f",
                         "spine_height_cm": 25.0, "spine_thickness_cm": 3.0, "replacement_cost": {"amount": None}, "used_value": {"amount": None}}],
              "items": [{"id": "I1", "category": "lamp", "description": "brass floor lamp"}],
              "room": {"floor_area_m2": 12.0, "wall_area_m2": 35.0, "length_m": 4, "width_m": 3, "height_m": 2.5},
              "totals": {"book_count": 3}, "meta": {"time_to_packet_s": 120}}
    d = tmp_path / "s"
    d.mkdir()
    (d / "claim_packet.json").write_text(json.dumps(packet))
    gt = tmp_path / "gt.xlsx"
    write_template(gt)
    wb = load_workbook(gt)
    wb["summary"]["B2"] = 3
    wb["books"].append(["g1", "", "", "Y", "The Hobbit", "Tolkien", ""])
    wb["books"].append(["g2", "", "", "Y", "Dune", "Herbert", ""])
    wb["books"].append(["g3", "", "", "N", "", "", ""])
    wb["spine_sample"].append(["g1", "The Hobbit", 21.0, 2.1, ""])
    wb["price_sample"].append(["The Hobbit", "new", "Shop", "u", "2026-10-09", 14.0, "USD", ""])
    wb["items"].append(["i1", "lamp", "floor lamp", ""])
    for r, v in ((2, 4.0), (3, 3.1), (4, 2.5)):
        wb["room"].cell(row=r, column=2, value=v)
    wb.save(gt)
    res = score(d, gt)
    b = res["bars"]
    assert b["book_count"]["pass"]
    assert b["title_identification"]["correct"] == 1 and b["title_identification"]["confidently_wrong"] == 1
    assert not b["title_identification"]["pass"]
    assert b["spine_dimensions"]["pass"]
    assert abs(b["book_prices"]["mape"] - 1 / 14) < 1e-6
    assert b["non_book_items"]["pass"]
    assert abs(b["floor_area"]["error"] - abs(12 - 12.4) / 12.4) < 1e-6


def test_server_requires_key():
    from fastapi.testclient import TestClient
    from library_claim.live import server
    c = TestClient(server.app)
    assert c.get("/").status_code == 403
    r = c.get(f"/?k={server.TOKEN}")
    assert r.status_code == 200 and server.TOKEN in r.text
    assert c.get("/health").json()["ok"]
    assert c.get(f"/sweeps/x/../../.env?k={server.TOKEN}").status_code == 404
